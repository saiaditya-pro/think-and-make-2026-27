"""
Row sources for the Google Sheets importers in ``apps.core.sheet_importers``.

Both sources expose the same tiny interface -- ``.sheetnames`` and
``.rows(tab)`` yielding ``{header: value}`` dicts -- so the importers don't
care whether the data came from a downloaded xlsx export or straight from
the live spreadsheet via the Sheets API.
"""

import json
from collections.abc import Iterator

from django.conf import settings
from openpyxl import load_workbook

SHEETS_API = "https://sheets.googleapis.com/v4/spreadsheets"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]


class SheetsConfigError(Exception):
    """Raised when the Google service-account settings are missing."""


def _normalize_header(value) -> str | None:
    # The master Users tab's first header is literally " Email".
    return str(value).strip() if value not in (None, "") else None


def _row_dicts(header_row, data_rows) -> Iterator[dict]:
    headers = [(_normalize_header(h), idx) for idx, h in enumerate(header_row)]
    for row in data_rows:
        if not row or row[0] in (None, ""):
            continue
        yield {name: (row[idx] if idx < len(row) else None) for name, idx in headers if name is not None}


class XlsxSource:
    def __init__(self, path):
        self._wb = load_workbook(path, data_only=True)

    @property
    def sheetnames(self) -> list[str]:
        return self._wb.sheetnames

    def rows(self, tab: str) -> Iterator[dict]:
        if tab not in self._wb.sheetnames:
            return iter(())
        sheet = self._wb[tab]
        all_rows = sheet.iter_rows(values_only=True)
        header_row = next(all_rows, ())
        return _row_dicts(header_row, all_rows)


def _authorized_session():
    from google.auth.transport.requests import AuthorizedSession
    from google.oauth2 import service_account

    if settings.GOOGLE_SERVICE_ACCOUNT_JSON:
        info = json.loads(settings.GOOGLE_SERVICE_ACCOUNT_JSON)
        credentials = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
    elif settings.GOOGLE_SERVICE_ACCOUNT_FILE:
        # Relative paths are relative to backend/, wherever the process was started.
        credentials = service_account.Credentials.from_service_account_file(
            settings.BASE_DIR / settings.GOOGLE_SERVICE_ACCOUNT_FILE, scopes=SCOPES
        )
    else:
        raise SheetsConfigError(
            "Set GOOGLE_SERVICE_ACCOUNT_FILE (or GOOGLE_SERVICE_ACCOUNT_JSON) in backend/.env."
        )
    return AuthorizedSession(credentials)


def _a1_tab(tab: str) -> str:
    return "'" + tab.replace("'", "''") + "'"


class GoogleSheetsSource:
    """
    Reads every tab of one spreadsheet with two API calls (metadata +
    values:batchGet), then serves rows from memory.

    Values come back UNFORMATTED (numbers stay numbers, checkboxes stay
    booleans) but dates come back as their formatted string, which
    ``sheet_importers._date`` knows how to parse.
    """

    def __init__(self, spreadsheet_id: str, session=None):
        self.spreadsheet_id = spreadsheet_id
        self._session = session or _authorized_session()
        self._values: dict[str, list[list]] = {}
        self._load()

    def _load(self):
        meta = self._session.get(
            f"{SHEETS_API}/{self.spreadsheet_id}", params={"fields": "sheets.properties.title"}
        )
        meta.raise_for_status()
        tabs = [s["properties"]["title"] for s in meta.json().get("sheets", [])]
        if not tabs:
            return

        resp = self._session.get(
            f"{SHEETS_API}/{self.spreadsheet_id}/values:batchGet",
            params=[("ranges", _a1_tab(t)) for t in tabs]
            + [("valueRenderOption", "UNFORMATTED_VALUE"), ("dateTimeRenderOption", "FORMATTED_STRING")],
        )
        resp.raise_for_status()
        for tab, value_range in zip(tabs, resp.json().get("valueRanges", [])):
            self._values[tab] = value_range.get("values", [])

    @property
    def sheetnames(self) -> list[str]:
        return list(self._values)

    def rows(self, tab: str) -> Iterator[dict]:
        values = self._values.get(tab) or []
        if not values:
            return iter(())
        return _row_dicts(values[0], values[1:])
